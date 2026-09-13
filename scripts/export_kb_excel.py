# -*- coding: utf-8 -*-
"""RAG 지식베이스 -> 엑셀"""
import csv, io, os, json, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
D = lambda *p: os.path.join(BASE, *p)

HF = PatternFill("solid", fgColor="1F4E79")
HFONT = Font(color="FFFFFF", bold=True, size=10)


def sheet(ws, path, widths, wrap_cols=()):
    rows = list(csv.reader(io.open(path, encoding="utf-8-sig")))
    for i, h in enumerate(rows[0], 1):
        c = ws.cell(1, i, h)
        c.fill = HF
        c.font = HFONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    for ri, r in enumerate(rows[1:], 2):
        for ci, v in enumerate(r, 1):
            cell = ws.cell(ri, ci, int(v) if v.isdigit() else v)
            if ci in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(rows[0])), len(rows))
    return len(rows) - 1


wb = openpyxl.Workbook()
n = {}
ws = wb.active
ws.title = "K-DST 발달기준"
n["kdst"] = sheet(ws, D("data", "kb", "kb_kdst.csv"), [9, 9, 13, 12, 6, 88], (6,))
n["cond"] = sheet(wb.create_sheet("질환-자폐 뇌성마비 다운"), D("data", "kb", "kb_conditions_all.csv"), [22, 26, 96], (3,))
n["lang"] = sheet(wb.create_sheet("질환-언어장애"), D("data", "kb", "kb_conditions.csv"), [20, 24, 96], (3,))
n["hear"] = sheet(wb.create_sheet("질환-유소아난청"), D("data", "kb", "kb_hearing.csv"), [18, 22, 96], (3,))
n["vid"] = sheet(wb.create_sheet("질환-시각 지적발달"), D("data", "kb", "kb_vision_id.csv"), [22, 24, 96], (3,))
n["ksg"] = sheet(wb.create_sheet("새싹과단비 발달특징"), D("data", "kb", "kb_ksied_guide.csv"), [12, 9, 9, 12, 88], (5,))
n["ksw"] = sheet(wb.create_sheet("새싹과단비 경고신호"), D("data", "kb", "kb_ksied_warning.csv"), [12, 9, 9, 92], (4,))
n["tic"] = sheet(wb.create_sheet("질환-틱장애"), D("data", "kb", "kb_tic.csv"), [22, 22, 96], (3,))
n["growth"] = sheet(wb.create_sheet("성장기준"), D("data", "kb", "kb_growth.csv"), [16, 18, 96], (3,))
n["area"] = sheet(wb.create_sheet("치료영역 설명"), D("data", "kb", "kb_therapy_areas.csv"), [16, 92, 16], (2,))
n["map"] = sheet(wb.create_sheet("영역 매핑"), D("data", "kb", "kb_mapping.csv"), [14, 16, 9, 60, 12], (4,))
n["policy"] = sheet(wb.create_sheet("바우처 제도"), D("data", "kb", "kb_policy.csv"), [16, 96, 22], (2,))

chunks = [json.loads(l) for l in io.open(D("data", "kb", "kb_chunks.jsonl"), encoding="utf-8")]
types = collections.Counter(c["type"] for c in chunks)
srcs = collections.Counter(c["source"] for c in chunks)

s = wb.create_sheet("개요", 0)
for col, w in (("A", 30), ("B", 12), ("C", 62)):
    s.column_dimensions[col].width = w


def put(r, a, b="", c="", bold=False):
    s.cell(r, 1, a).font = Font(bold=bold, size=11 if bold else 10)
    if b != "":
        s.cell(r, 2, b)
    if c != "":
        s.cell(r, 3, c).alignment = Alignment(wrap_text=True, vertical="top")


put(1, "안아수달 RAG 지식베이스", bold=True)
put(2, "총 청크", len(chunks), "kb_chunks.jsonl — 임베딩용, 출처 메타 포함")
r = 4
put(r, "시트", "건수", "내용", bold=True)
r += 1
for lab, k, desc in (
    ("K-DST 발달기준", "kdst", "질병관리청 영유아 발달선별검사 16개 연령구간 6영역"),
    ("새싹과단비 발달특징", "ksg", "삼성복지재단·육아정책연구소 11개 연령구간 5영역"),
    ("새싹과단비 경고신호", "ksw", "연령별 발달 경고신호"),
    ("질환-자폐 뇌성마비 다운", "cond", "자폐스펙트럼장애 · 뇌성마비 · 다운증후군"),
    ("질환-언어장애", "lang", "아동 언어장애 진단과 치료, 경고신호, 검사도구"),
    ("질환-유소아난청", "hear", "난청 정도·증상·원인·검사·1-3-6 타임라인·고위험인자"),
    ("질환-시각 지적발달", "vid", "굴절이상·약시 / 지적발달장애 진단기준·심각도"),
    ("질환-틱장애", "tic", "틱 종류·진단분류·경과·치료 판단기준·가족 대처"),
    ("성장기준", "growth", "정상 성장 속도, 미숙아 교정연령, 머리둘레 경고신호"),
    ("치료영역 설명", "area", "법정 치료영역 10종을 부모 언어로 설명"),
    ("영역 매핑", "map", "K-DST 영역 → 법정 치료영역 + 우선순위"),
    ("바우처 제도", "policy", "지원대상, 소득기준, 지원금액, 신청 절차"),
):
    put(r, lab, n[k], desc)
    r += 1

r += 1
put(r, "청크 유형별", bold=True)
r += 1
for t, c in types.most_common():
    put(r, t, c)
    r += 1

r += 1
put(r, "출처별", bold=True)
r += 1
for src, c in srcs.most_common():
    put(r, src[:28], c, src)
    r += 1

r += 1
put(r, "검색 설계 권고", bold=True)
r += 1
for a, b in (
    ("1. 연령 메타 필터 먼저", "월령으로 K-DST 구간을 좁히면 736 → 40~48문항. 검색 공간 1/15"),
    ("2. 미숙아 교정연령", "37주 이전 출생아는 만 2세까지 출생예정일 기준으로 월령 계산"),
    ("3. 유사도 임계값", "90퍼센트는 거의 안 걸림. 질문셋으로 정답 문서의 유사도 분포를 측정해 결정"),
    ("4. temperature", "생성 단계 파라미터. 근거 기반 답변이므로 0.2 내외로 낮게 고정"),
    ("5. 진단 표현 금지", "'OO 의심' 대신 '확인해볼 영역'으로 표현하고 진료과를 안내"),
):
    put(r, a, "", b)
    r += 1

r += 1
put(r, "주의", bold=True)
r += 1
put(r, "매핑 검증 상태", "", "영역 매핑 15건 중 10건은 국가건강정보포털 치료법 기술에 근거(evidence=source), 5건은 자체 초안(draft). draft 항목은 전문가 자문 필요")
r += 1
put(r, "면책", "", "국가건강정보포털 콘텐츠는 참고사항이며 법적 책임이 없음을 명시. 서비스에서도 동일한 면책 문구 필요")

out = D("안아수달_RAG지식베이스.xlsx")
wb.save(out)
print("저장: %s" % os.path.basename(out))
print("총 청크 %d / 시트 %d" % (len(chunks), len(wb.sheetnames)))
