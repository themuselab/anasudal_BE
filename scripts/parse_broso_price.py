# -*- coding: utf-8 -*-
"""broso 시·도별 발달재활서비스 가격 공시(서울) 파싱 + 마스터 목록 매칭.
출처: 중앙장애아동·발달장애인지원센터 정책·법령·통계 > 2026년 시·도별 발달재활서비스 가격 공시
치료영역(제공영역) + 회기당 단가 + 방문 가능 여부를 공공데이터로 확보하는 경로.
"""
import openpyxl, csv, io, os, re, collections

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
D = lambda *p: os.path.join(BASE, *p)

wb = openpyxl.load_workbook(D("data","raw","seoul_price.xlsx"), read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
hdr = [str(c).replace("\n", "") if c else "" for c in rows[0]]
I = {k: i for i, k in enumerate(hdr)}
C_GU, C_NM, C_BIZ = I["시·군·구"], I["제공기관명"], I["사업자등록번호"]
C_WAY = [i for i, k in enumerate(hdr) if "제공 방식" in k][0]
C_AREA, C_P26 = I["제공영역"], I["26년단가(원)"]

recs = [r for r in rows[1:] if r[C_NM]]

# 기관 단위로 접기
agg = {}
for r in recs:
    key = (str(r[C_BIZ]).strip(), str(r[C_NM]).strip())
    a = agg.setdefault(key, {"name": str(r[C_NM]).strip(), "biz_no": str(r[C_BIZ]).strip(),
                             "sigungu": str(r[C_GU]).strip(), "areas": set(),
                             "visit": False, "onsite": False, "prices": []})
    if r[C_AREA]: a["areas"].add(str(r[C_AREA]).strip())
    way = str(r[C_WAY] or "")
    if "방문" in way: a["visit"] = True
    if "기관" in way: a["onsite"] = True
    if isinstance(r[C_P26], (int, float)): a["prices"].append(int(r[C_P26]))

out = []
for a in agg.values():
    p = sorted(a["prices"])
    out.append({
        "name": a["name"], "biz_no": a["biz_no"], "sigungu": a["sigungu"],
        "areas": " / ".join(sorted(a["areas"])), "area_count": len(a["areas"]),
        "price_min": p[0] if p else "", "price_max": p[-1] if p else "",
        "onsite": "Y" if a["onsite"] else "", "visit": "Y" if a["visit"] else "",
        "source": "broso_2026_price",
    })
out.sort(key=lambda r: (r["sigungu"], r["name"]))

with io.open(D("data","seoul_broso_price.csv"), "w", encoding="utf-8-sig", newline="") as fp:
    w = csv.DictWriter(fp, fieldnames=list(out[0].keys())); w.writeheader(); w.writerows(out)

# 마스터(377)와 매칭
norm = lambda s: re.sub(r"[\s()\[\]㈜·\.]", "", s or "")
master = list(csv.DictReader(io.open(D("data","seoul_master.csv"), encoding="utf-8-sig")))
midx = {norm(m["name"]): m for m in master}
hit = sum(1 for r in out if norm(r["name"]) in midx)

o = io.open(D("data","broso_price_report.txt"), "w", encoding="utf-8"); w = lambda *a: o.write(" ".join(str(z) for z in a) + "\n")
w("broso 2026 서울 가격 공시")
w("  원본 행수(기관×영역×방식) : %d" % len(recs))
w("  고유 기관 수               : %d" % len(out))
w("  마스터 377건과 이름 일치   : %d (%.0f%%)" % (hit, 100.0*hit/len(out)))
w("")
w("=== 제공영역 분포 (기관 수) ===")
ca = collections.Counter(a for r in out for a in r["areas"].split(" / ") if a)
for k, v in ca.most_common(): w("   %-12s %4d  (%.0f%%)" % (k, v, 100.0*v/len(out)))
w("")
w("=== 기관당 제공영역 개수 ===")
for k, v in sorted(collections.Counter(r["area_count"] for r in out).items()): w("   %d개 영역: %d곳" % (k, v))
w("")
allp = sorted(p for r in out if r["price_min"] for p in (r["price_min"], r["price_max"]))
w("=== 26년 단가 ===")
w("   최소 %s / 중앙 %s / 최대 %s원" % (f"{allp[0]:,}", f"{allp[len(allp)//2]:,}", f"{allp[-1]:,}"))
w("   방문 서비스 제공: %d곳 (%.0f%%)" % (sum(1 for r in out if r["visit"]), 100.0*sum(1 for r in out if r["visit"])/len(out)))
w("")
w("=== 자치구별 ===")
for k, v in collections.Counter(r["sigungu"] for r in out).most_common()[:8]: w("   %-8s %d" % (k, v))
w("")
w("=== 샘플 ===")
for r in out[:4]: w("   %s | %s | %s원~%s원 | %s" % (r["name"][:18], r["areas"][:60], f'{r["price_min"]:,}', f'{r["price_max"]:,}', "방문가능" if r["visit"] else "기관내"))
o.close()
print("ok")
