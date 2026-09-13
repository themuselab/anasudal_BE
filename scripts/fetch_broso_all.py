# -*- coding: utf-8 -*-
"""broso '2026년 시·도별 발달재활서비스 가격 공시' 17개 시·도 xlsx 전량 수집 -> 전국 CSV"""
import re, io, os, csv, json, subprocess, collections, urllib.parse
import openpyxl

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
D = lambda *p: os.path.join(BASE, *p)
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120"
RAW = D("data", "raw", "broso")
os.makedirs(RAW, exist_ok=True)

h = open(D("data", "raw", "broso_post.html"), "rb").read().decode("utf-8", "replace")
files = []
for m in re.finditer(r"data-data='(\{.*?\})'", h, re.S):
    try:
        d = json.loads(re.sub(r"\s+", " ", m.group(1)))
    except Exception:
        continue
    if d.get("orignlFileNm") and d.get("streFileNm"):
        files.append(d)
print("첨부파일 %d개" % len(files), flush=True)

for f in files:
    nm = f["orignlFileNm"]
    dst = os.path.join(RAW, re.sub(r"[^\w가-힣.]", "_", nm))
    if not (os.path.exists(dst) and os.path.getsize(dst) > 5000):
        url = ("https://www.broso.or.kr/hp/cmm/downFile.do?fileSn=%s&atchFileId=%s"
               "&streFileNm=%s&isTemp=&orignFileNm=%s&mode=down"
               % (f["fileSn"], f["atchFileId"], f["streFileNm"], urllib.parse.quote(nm)))
        subprocess.run(["curl", "-sL", "-m", "90", "-k", "-A", UA,
                        "-H", "Referer: https://www.broso.or.kr/hp/bbs/ViewCmpPosNttDtl.do",
                        url, "-o", dst])
    print("  %-26s %8d B" % (nm, os.path.getsize(dst)), flush=True)

out, bad = [], []
for fn in sorted(os.listdir(RAW)):
    p = os.path.join(RAW, fn)
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception as e:
        bad.append((fn, str(e)[:40])); continue
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).replace("\n", "").strip() if c else "" for c in rows[0]]
    def col(*names):
        for n in names:
            for i, k in enumerate(hdr):
                if n in k:
                    return i
        return None
    c_sido = col("시·도", "시도")
    c_gu   = col("시·군·구", "시군구", "군구")
    c_nm   = col("제공기관명", "기관명")
    c_biz  = col("사업자등록번호")
    c_way  = col("제공 방식", "제공방식")
    c_area = col("제공영역", "영역")
    c_p    = col("26년단가", "2026년 단가", "26년 단가") or col("25년단가", "2025년 단가")
    if c_gu is None or c_nm is None or c_area is None:
        bad.append((fn, "컬럼불일치 %s" % hdr[:6])); continue
    for r in rows[1:]:
        if not r[c_nm]:
            continue
        out.append({
            "sido":    str(r[c_sido]).strip() if c_sido is not None and r[c_sido] else "",
            "sigungu": str(r[c_gu]).strip() if r[c_gu] else "",
            "name":    str(r[c_nm]).strip(),
            "biz_no":  str(r[c_biz]).strip() if c_biz is not None and r[c_biz] else "",
            "way":     str(r[c_way]).strip() if c_way is not None and r[c_way] else "",
            "area":    str(r[c_area]).strip(),
            "price":   r[c_p] if c_p is not None and isinstance(r[c_p], (int, float)) else "",
            "src_file": fn,
        })

with io.open(D("data", "broso_price_nationwide_raw.csv"), "w", encoding="utf-8-sig", newline="") as fp:
    w = csv.DictWriter(fp, fieldnames=list(out[0].keys())); w.writeheader(); w.writerows(out)

agg = {}
for r in out:
    k = (r["biz_no"] or r["name"], r["name"], r["sigungu"])
    a = agg.setdefault(k, {"sido": r["sido"], "sigungu": r["sigungu"], "name": r["name"],
                           "biz_no": r["biz_no"], "areas": set(), "visit": False, "prices": []})
    a["areas"].add(r["area"])
    if "방문" in r["way"]:
        a["visit"] = True
    if r["price"]:
        a["prices"].append(int(r["price"]))

fin = []
for a in agg.values():
    p = sorted(a["prices"])
    fin.append({"sido": a["sido"], "sigungu": a["sigungu"], "name": a["name"], "biz_no": a["biz_no"],
                "areas": " / ".join(sorted(a["areas"])), "area_count": len(a["areas"]),
                "price_min": p[0] if p else "", "price_max": p[-1] if p else "",
                "visit": "Y" if a["visit"] else ""})
fin.sort(key=lambda r: (r["sido"], r["sigungu"], r["name"]))
with io.open(D("data", "broso_nationwide.csv"), "w", encoding="utf-8-sig", newline="") as fp:
    w = csv.DictWriter(fp, fieldnames=list(fin[0].keys())); w.writeheader(); w.writerows(fin)

print("\n원본 행 %d / 고유 기관 %d" % (len(out), len(fin)))
print("파싱 실패: %s" % bad)
print("시도 %d개" % len({r["sido"] for r in fin if r["sido"]}))
