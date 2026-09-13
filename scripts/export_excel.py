# -*- coding: utf-8 -*-
"""전국 마스터 -> 엑셀 (기관 시트 + 요약 시트 + 영역별 원본 시트)"""
import csv, io, os, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
D = lambda *p: os.path.join(BASE, *p)
rows = list(csv.DictReader(io.open(D("data", "nationwide_master.csv"), encoding="utf-8-sig")))

COLS = [
    ("사업자등록번호","biz_no",15), ("기관명","name",30), ("시도","sido",14), ("시군구","sigungu",14),
    ("도로명주소","address",44), ("전화번호","tel",15), ("이메일","email",26),
    ("치료영역","areas",46), ("영역수","area_count",7),
    ("단가최저","price_min",10), ("단가최고","price_max",10), ("방문서비스","visit",9),
    ("자체홈페이지","own_homepage",34), ("카카오플레이스","homepage",36),
    ("업종분류","kakao_category",30), ("프로그램내용","programs_local",26),
    ("위도","lat",13), ("경도","lon",13), ("매칭신뢰도","match_conf",18), ("출처","src",16),
]
NULL_COLS = [("특화분야(세부)","specialty_extra",14),("운영시간","operating_hours",12),
             ("대기여부","waiting",10),("중증도수용","severity",11),
             ("대상연령","age_range",10),("치료사수","therapist_count",10)]

wb = openpyxl.Workbook()
hdr_fill = PatternFill("solid", fgColor="1F4E79")
null_fill = PatternFill("solid", fgColor="9E9E9E")
hdr_font = Font(color="FFFFFF", bold=True, size=10)
thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# ── 시트1: 기관 마스터
ws = wb.active; ws.title = "기관마스터"
allcols = COLS + NULL_COLS
for i, (label, _, wdt) in enumerate(allcols, 1):
    c = ws.cell(1, i, label)
    c.fill = null_fill if i > len(COLS) else hdr_fill
    c.font = hdr_font; c.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[get_column_letter(i)].width = wdt
for r_i, r in enumerate(rows, 2):
    for c_i, (_, key, _) in enumerate(allcols, 1):
        v = r.get(key, "")
        if key in ("price_min", "price_max", "area_count") and str(v).isdigit():
            v = int(v)
        cell = ws.cell(r_i, c_i, v)
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=(key == "areas"))
ws.freeze_panes = "C2"
ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(allcols)), len(rows) + 1)

# ── 시트2: 요약
s = wb.create_sheet("요약")
n = len(rows)
def put(r, a, b="", c="", bold=False):
    s.cell(r, 1, a).font = Font(bold=bold or (b == ""), size=11 if b == "" else 10)
    if b != "": s.cell(r, 2, b)
    if c != "": s.cell(r, 3, c)
s.column_dimensions["A"].width = 26; s.column_dimensions["B"].width = 14; s.column_dimensions["C"].width = 46
put(1, "전국 발달재활 기관 마스터", bold=True)
put(2, "총 기관 수", n, "%d개 시도" % len({r["sido"] for r in rows if r["sido"]}))
r_ = 4
put(r_, "필드", "충족률", "출처"); r_ += 1
SRC = {"biz_no":"broso 가격공시","areas":"broso 가격공시","price_min":"broso 가격공시",
       "visit":"broso 가격공시","address":"복지부+지자체+사회보장정보원","tel":"사회보장정보원+카카오+지자체",
       "email":"사회보장정보원","own_homepage":"카카오 플레이스","homepage":"카카오 로컬",
       "lat":"카카오 로컬","kakao_category":"카카오 로컬","programs_local":"지자체 공공데이터"}
for label, key, _ in COLS:
    if key in ("name","sido","sigungu","area_count","price_max","match_conf","src"): continue
    c = sum(1 for x in rows if x.get(key))
    put(r_, label, "%.1f%%" % (100.0*c/n), SRC.get(key, "")); r_ += 1
r_ += 1
put(r_, "공공·카카오 모두 없음 (null)", bold=True); r_ += 1
for label, key, _ in NULL_COLS:
    put(r_, label, "0.0%", "전화조사 또는 기관입력 필요"); r_ += 1
r_ += 1
put(r_, "매칭 신뢰도", bold=True); r_ += 1
for k, v in collections.Counter(x["match_conf"] for x in rows).most_common():
    put(r_, k or "(미처리)", v, "%.1f%%" % (100.0*v/n)); r_ += 1

# ── 시트3: 치료영역별 원본(기관×영역×방식)
raw = list(csv.DictReader(io.open(D("data", "broso_price_nationwide_raw.csv"), encoding="utf-8-sig")))
s3 = wb.create_sheet("치료영역_단가원본")
h3 = [("시도","sido",14),("시군구","sigungu",14),("기관명","name",30),("사업자등록번호","biz_no",15),
      ("제공방식","way",10),("제공영역","area",14),("단가(원)","price",11)]
for i,(label,_,wdt) in enumerate(h3,1):
    c=s3.cell(1,i,label); c.fill=hdr_fill; c.font=hdr_font
    c.alignment=Alignment(horizontal="center"); s3.column_dimensions[get_column_letter(i)].width=wdt
for r_i, r in enumerate(raw, 2):
    for c_i,(_,key,_) in enumerate(h3,1):
        v=r.get(key,"")
        if key=="price" and str(v).isdigit(): v=int(v)
        s3.cell(r_i,c_i,v)
s3.freeze_panes="A2"; s3.auto_filter.ref="A1:G%d"%(len(raw)+1)

out = D("안아수달_전국기관DB.xlsx")
wb.save(out)
print("저장: %s" % out)
print("  기관마스터 %d행 x %d열 / 치료영역원본 %d행" % (len(rows), len(allcols), len(raw)))
